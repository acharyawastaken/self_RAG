import hashlib
import json
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
import filetype
import tiktoken
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from markdown_it import MarkdownIt

from backend.app.config import get_settings

settings = get_settings()

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class IngestionError(Exception):
    pass


def compute_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def validate_file(filename: str, content: bytes) -> str:
    path = Path(filename)
    ext = path.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"Unsupported file extension '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    max_bytes = settings.max_upload_mb * 1024 * 1024
    if len(content) > max_bytes:
        raise IngestionError(f"File size exceeds maximum limit of {settings.max_upload_mb} MB")

    if len(content) == 0:
        raise IngestionError("Uploaded file is empty (0 bytes)")

    # MIME / magic bytes verification
    guess = filetype.guess(content)
    if ext == ".pdf":
        if not content.startswith(b"%PDF"):
            raise IngestionError("Invalid PDF file: Missing %PDF header")
    elif ext == ".docx":
        # docx is a zip file containing word/document.xml
        if not (guess and guess.extension in ["zip", "docx"]):
            # Also check zip magic bytes b"PK\x03\x04"
            if not content.startswith(b"PK\x03\x04"):
                raise IngestionError("Invalid DOCX file: File format is not a valid OpenXML Word document")
    elif ext in [".txt", ".md"]:
        try:
            content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                import charset_normalizer
                match = charset_normalizer.from_bytes(content).best()
                if not match:
                    raise IngestionError(f"Cannot decode {ext} file as text")
            except Exception as e:
                raise IngestionError(f"Encoding error in text file: {str(e)}")

    return ext


def extract_pdf(content: bytes) -> Tuple[List[Dict[str, Any]], int]:
    doc = fitz.open(stream=content, filetype="pdf")
    page_count = len(doc)
    sections: List[Dict[str, Any]] = []

    for i in range(page_count):
        page = doc[i]
        text = page.get_text("text").strip()
        if text:
            sections.append({
                "text": text,
                "page": i + 1,
                "section_title": f"Page {i + 1}"
            })
    doc.close()
    return sections, page_count


def extract_docx(content: bytes) -> Tuple[List[Dict[str, Any]], int]:
    import io
    doc = DocxDocument(io.BytesIO(content))
    sections: List[Dict[str, Any]] = []
    current_title = "Introduction"
    current_paras: List[str] = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style.name.startswith("Heading"):
            if current_paras:
                sections.append({
                    "text": "\n".join(current_paras),
                    "page": None,
                    "section_title": current_title
                })
                current_paras = []
            current_title = txt
        else:
            current_paras.append(txt)

    if current_paras:
        sections.append({
            "text": "\n".join(current_paras),
            "page": None,
            "section_title": current_title
        })

    # Also extract any table content
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_txt:
                table_texts.append(row_txt)
    if table_texts:
        sections.append({
            "text": "\n".join(table_texts),
            "page": None,
            "section_title": "Tables"
        })

    return sections, 1


def extract_markdown(content: bytes) -> Tuple[List[Dict[str, Any]], int]:
    text = content.decode("utf-8", errors="replace")
    md = MarkdownIt()
    tokens = md.parse(text)

    sections: List[Dict[str, Any]] = []
    current_title = "Overview"
    current_lines: List[str] = []

    for i, token in enumerate(tokens):
        if token.type == "heading_open":
            if current_lines:
                sections.append({
                    "text": "\n".join(current_lines),
                    "page": None,
                    "section_title": current_title
                })
                current_lines = []
            # Next token is inline with title
            if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                current_title = tokens[i + 1].content.strip()
        elif token.type in ["paragraph_open", "fence", "code_block"]:
            if i + 1 < len(tokens) and tokens[i + 1].content:
                current_lines.append(tokens[i + 1].content.strip())

    if current_lines:
        sections.append({
            "text": "\n".join(current_lines),
            "page": None,
            "section_title": current_title
        })

    # Fallback if markdown parser returned no paragraphs (e.g. raw text)
    if not sections and text.strip():
        sections.append({
            "text": text.strip(),
            "page": None,
            "section_title": "Document Content"
        })

    return sections, 1


def extract_txt(content: bytes) -> Tuple[List[Dict[str, Any]], int]:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        import charset_normalizer
        match = charset_normalizer.from_bytes(content).best()
        text = str(match) if match else content.decode("utf-8", errors="replace")

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    sections: List[Dict[str, Any]] = []
    for idx, p in enumerate(paragraphs):
        sections.append({
            "text": p,
            "page": None,
            "section_title": f"Section {idx + 1}"
        })
    if not sections and text.strip():
        sections.append({
            "text": text.strip(),
            "page": None,
            "section_title": "Full Text"
        })
    return sections, 1


def extract_document_sections(file_type: str, content: bytes) -> Tuple[List[Dict[str, Any]], Optional[int]]:
    ext = file_type.lower()
    if ext == ".pdf":
        return extract_pdf(content)
    elif ext == ".docx":
        return extract_docx(content)
    elif ext == ".md":
        return extract_markdown(content)
    elif ext == ".txt":
        return extract_txt(content)
    else:
        raise IngestionError(f"Unsupported file type: {file_type}")


_tokenizer = None


def get_tokenizer():
    global _tokenizer
    if _tokenizer is None:
        _tokenizer = tiktoken.get_encoding("cl100k_base")
    return _tokenizer


def chunk_sections(
    document_id: str,
    sections: List[Dict[str, Any]],
    chunk_size: int = 350,
    overlap: int = 60,
) -> List[Dict[str, Any]]:
    tokenizer = get_tokenizer()
    chunks: List[Dict[str, Any]] = []
    chunk_index = 0

    for section in sections:
        text = section["text"].strip()
        if not text:
            continue

        tokens = tokenizer.encode(text)
        total_tokens = len(tokens)

        if total_tokens <= chunk_size:
            chunk_id = f"{document_id}:{chunk_index}"
            chunks.append({
                "chunk_id": chunk_id,
                "document_id": document_id,
                "chunk_index": chunk_index,
                "text_content": text,
                "page_start": section.get("page"),
                "page_end": section.get("page"),
                "section_title": section.get("section_title"),
                "token_count": total_tokens,
                "metadata_json": json.dumps({
                    "document_id": document_id,
                    "section_title": section.get("section_title"),
                    "page": section.get("page"),
                })
            })
            chunk_index += 1
        else:
            # Slide window across tokens
            start = 0
            while start < total_tokens:
                end = min(start + chunk_size, total_tokens)
                sub_tokens = tokens[start:end]
                sub_text = tokenizer.decode(sub_tokens)

                chunk_id = f"{document_id}:{chunk_index}"
                chunks.append({
                    "chunk_id": chunk_id,
                    "document_id": document_id,
                    "chunk_index": chunk_index,
                    "text_content": sub_text,
                    "page_start": section.get("page"),
                    "page_end": section.get("page"),
                    "section_title": section.get("section_title"),
                    "token_count": len(sub_tokens),
                    "metadata_json": json.dumps({
                        "document_id": document_id,
                        "section_title": section.get("section_title"),
                        "page": section.get("page"),
                        "token_offset": start,
                    })
                })
                chunk_index += 1
                if end == total_tokens:
                    break
                start += (chunk_size - overlap)

    return chunks
